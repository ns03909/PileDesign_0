# -*- coding: utf-8 -*-
"""PileDesign 2026-04〜09 の改良点スライド (16:9) を生成する。"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

REPO = r"c:\Users\keisu\source\repos\PileDesign_0"
SHOTS = os.path.join(REPO, "Graphics_r1", "Help", "images", "screenshots")
OUT = os.path.join(REPO, "Docs", "PileDesign_変更点_2026-04_2026-09.pptx")

NAVY = RGBColor(0x1F, 0x35, 0x5E)
ACCENT = RGBColor(0x2E, 0x74, 0xB5)
GRAY = RGBColor(0x59, 0x59, 0x59)
LIGHT = RGBColor(0xEE, 0xF3, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x26, 0x26, 0x26)
FONT = "Yu Gothic UI"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def tb(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, text, size=14, bold=False, color=INK,
         space_before=0, space_after=4, first=False, line=1.25):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT
    return p


def rect(slide, x, y, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    return sh


def header(slide, no, title, sub=None):
    rect(slide, 0, 0, SW, Inches(1.02), NAVY)
    rect(slide, 0, Inches(1.02), SW, Inches(0.05), ACCENT)
    tf = tb(slide, Inches(0.62), Inches(0.14), Inches(12.0), Inches(0.78),
            anchor=MSO_ANCHOR.MIDDLE)
    para(tf, title, size=26, bold=True, color=WHITE, first=True, space_after=0)
    if sub:
        para(tf, sub, size=12.5, color=RGBColor(0xC5, 0xD6, 0xEA), space_after=0)
    ptf = tb(slide, SW - Inches(1.1), SH - Inches(0.5), Inches(0.6), Inches(0.3))
    p = para(ptf, str(no), size=11, color=GRAY, first=True)
    p.alignment = PP_ALIGN.RIGHT


def bullets(slide, x, y, w, items, size=13.5, gap=9):
    """items: (レベル, 文字列) のリスト。レベル 0 は小見出し。"""
    tf = tb(slide, x, y, w, Inches(0.4))
    for i, (lv, text) in enumerate(items):
        if lv == 0:
            para(tf, text, size=size + 1.5, bold=True, color=NAVY,
                 space_before=(0 if i == 0 else gap), space_after=3,
                 first=(i == 0))
        else:
            para(tf, "・" + text, size=size, color=INK, space_after=3,
                 first=(i == 0), line=1.22)
    return tf


def caption(slide, x, y, w, text, align=PP_ALIGN.CENTER):
    tf = tb(slide, x, y, w, Inches(0.3))
    p = para(tf, text, size=10.5, color=GRAY, first=True)
    p.alignment = align


def picture(slide, name, x, y, w):
    pic = slide.shapes.add_picture(os.path.join(SHOTS, name), x, y, width=w)
    pic.line.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
    pic.line.width = Pt(0.75)
    return pic


# ── 1. 表紙 ──────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, Inches(3.62), SW, Inches(0.055), ACCENT)

tf = tb(s, Inches(1.1), Inches(2.0), Inches(11.2), Inches(1.5))
para(tf, "杭基礎検討プログラム", size=22, color=RGBColor(0xA9, 0xC4, 0xE4),
     first=True, space_after=8)
para(tf, "この 4 か月半で変わったこと", size=44, bold=True, color=WHITE, space_after=0)

tf = tb(s, Inches(1.1), Inches(3.95), Inches(11.2), Inches(1.6))
para(tf, "機能の追加と、使い勝手の改善を中心に", size=17,
     color=RGBColor(0xC5, 0xD6, 0xEA), first=True, space_after=16)
para(tf, "2026 年 4 月 21 日 〜 9 月 6 日   /   v1.0.4-beta → v1.0.30-beta（22 版）",
     size=14, color=RGBColor(0x9F, 0xB8, 0xDB), space_after=0)

# ── 2. 全体像 ────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
header(s, 2, "全体像", "4 か月半で 388 コミット・22 版。次の 3 つを進めました。")

pillars = [
    ("1", "検討できる対象を広げた",
     ["杭工法・杭断面・設計法を追加。", "メーカー製品を選ぶだけで諸元が入る。"]),
    ("2", "結果が読めるようになった",
     ["検定を数値（検定比）で表示。", "図と計算書に、そのまま出る。"]),
    ("3", "日々の操作を軽くした",
     ["画面を刷新し、待ち時間を減らし、", "入力のやり直しを減らした。"]),
]
cw, gapx = Inches(3.85), Inches(0.42)
x0 = (SW - (cw * 3 + gapx * 2)) / 2
for i, (num, title, body) in enumerate(pillars):
    x = x0 + i * (cw + gapx)
    rect(s, x, Inches(1.55), cw, Inches(2.55), LIGHT)
    rect(s, x, Inches(1.55), cw, Inches(0.075), ACCENT)
    t = tb(s, x + Inches(0.3), Inches(1.85), cw - Inches(0.6), Inches(0.5))
    para(t, num, size=30, bold=True, color=ACCENT, first=True, space_after=2)
    para(t, title, size=17, bold=True, color=NAVY, space_after=8)
    for ln in body:
        para(t, ln, size=13, color=RGBColor(0x3A, 0x3A, 0x3A), space_after=1, line=1.3)

rect(s, Inches(0.72), Inches(4.55), SW - Inches(1.44), Inches(1.42),
     WHITE, line=RGBColor(0xC9, 0xD5, 0xE4))
stats = [
    ("22 版", "1.0.4-beta → 1.0.30-beta"),
    ("388", "コミット"),
    ("約 1,500 件", "自動テスト（4 月は 295 件）"),
    ("45 / 45", "文献値と一致した照合項目"),
]
bw = (SW - Inches(1.44)) / 4
for i, (big, small) in enumerate(stats):
    x = Inches(0.72) + i * bw
    t = tb(s, x, Inches(4.78), bw, Inches(1.0), anchor=MSO_ANCHOR.MIDDLE)
    p = para(t, big, size=25, bold=True, color=NAVY, first=True, space_after=3)
    p.alignment = PP_ALIGN.CENTER
    p = para(t, small, size=11.5, color=GRAY, space_after=0)
    p.alignment = PP_ALIGN.CENTER

t = tb(s, Inches(0.72), Inches(6.25), SW - Inches(1.44), Inches(0.5))
para(t, "※ 不具合の修正もあわせて進めましたが、本資料では機能と使い勝手を中心に説明します"
        "（全件は CHANGELOG.md に記録）。",
     size=11.5, color=GRAY, first=True)

# ── 3. 対象範囲 ──────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
header(s, 3, "① 検討できる対象を広げた",
       "扱える工法・断面・設計法が増え、製品を選ぶだけで諸元が入ります。")

left = [
    (0, "杭工法"),
    (1, "キャプリングパイル工法（杭頭の半剛接合）／鋼管杭＋鉄筋定着工法"),
    (1, "Smart-MAGNUM 工法（ジャパンパイル）の支持力"),
    (1, "Hybrid ニーディング工法（三谷セキサン）の支持力"),
    (1, "KCTB（TB 工法）による場所打ち鋼管コンクリート杭の設計法"),
    (0, "杭断面と製品ライブラリ"),
    (1, "PHC 節杭・PRC 節杭・BF.S パイル（頭部厚型節付き杭）を追加"),
    (1, "メーカー製品 1,469 行を収録"),
    (1, "　JP-NPH／JP-NPRC／BF.S／MS-hi105／Hi-SC105／DAM105"),
    (1, "製品を選べば諸元が入り、規格の違いは諸元表が知らせます"),
]
right = [
    (0, "設計法の選択肢"),
    (1, "2025 年版 技術基準解説書 付録 1-3 に準拠するオプション"),
    (1, "　告示 1113（第 8）の許容圧縮・許容せん断へ切替"),
    (1, "材料のモデル化を基本設定で選択（全 9 項目）"),
    (1, "　Ec の ξ＝1.0／圧縮 0.85Fc／鉄筋・鋼管の 1.1F バイリニア ほか"),
    (1, "許容時 N-M を「ファイバーモデル」と「単純累加」から選択"),
    (0, "解析"),
    (1, "ファイバーモデル M-φ をコンクリート系の全杭種へ拡大"),
    (1, "地盤水平変位に第 3 の算定法「応答スペクトル法」を追加"),
    (1, "地盤 p-y の非線形性を 3 段階から荷重ケースごとに選択"),
    (1, "基準水平地盤反力係数 kh0 を土層ごとに手入力で上書き可能に"),
]
bullets(s, Inches(0.72), Inches(1.55), Inches(5.9), left, size=14.5, gap=14)
bullets(s, Inches(7.05), Inches(1.55), Inches(5.6), right, size=14.5, gap=14)

rect(s, Inches(0.72), Inches(6.12), SW - Inches(1.44), Inches(0.8), LIGHT)
t = tb(s, Inches(1.05), Inches(6.22), SW - Inches(2.1), Inches(0.6), anchor=MSO_ANCHOR.MIDDLE)
para(t, "工法と製品を選べば、諸元・支持力・耐力の算定までそのまま通ります",
     size=14.5, bold=True, color=NAVY, first=True)

# ── 4. 検定 ──────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
header(s, 4, "② 結果が読めるようになった",
       "「OK / NG」だけだった検定を、数値と色で示すようにしました。")

items = [
    (0, "検定を数値に"),
    (1, "検定比と支配ケースを表示。余裕がどれだけ有るかが分かります"),
    (1, "検定項目を拡張"),
    (1, "　鉛直支持力（押込み・引抜き）／杭体のせん断"),
    (1, "　杭頭 2 点間の変形角／長期（常時）の曲げ"),
    (1, "支持力の検定を計算書 (docx) にも出力"),
    (0, "見せ方"),
    (1, "解析結果ダッシュボードを「検定の総括」に作り替え"),
    (1, "杭配置図を検定比で色分け。杭にマウスを近づけると、\n　  その部位ごとの検定比を表示"),
    (1, "収束しなかったケースは、OK でも NG でもなく「未収束」と表示"),
]
bullets(s, Inches(0.72), Inches(1.55), Inches(5.4), items, size=14.5, gap=14)

picture(s, "plan_evaluation.png", Inches(6.45), Inches(1.52), Inches(6.2))
caption(s, Inches(6.45), Inches(4.93), Inches(6.2), "平面図を検定比で色分けし、ホバーで内訳を表示")

picture(s, "dashboard.png", Inches(6.45), Inches(5.33), Inches(3.95))
caption(s, Inches(10.6), Inches(6.1), Inches(2.2), "検定の総括\nダッシュボード", align=PP_ALIGN.LEFT)

# ── 5. 使い勝手 ──────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
header(s, 5, "③ 日々の操作を軽くした", "画面・入力・計算書の、手数と待ち時間を減らしました。")

col1 = [
    (0, "画面と操作"),
    (1, "UI を全面刷新（リボン、ファイルメニューに計算例 11 例題）"),
    (1, "全 18 ウィンドウでキーボード操作を統一。ショートカット一覧を常設"),
    (1, "コマンドパレット (Ctrl+Shift+P)、編集履歴パネル、ヘルプチャット"),
    (1, "押せないボタン・効かないキーは、理由をその場で表示"),
    (0, "入力のやり直しを減らす"),
    (1, "入力を変えても解析結果が消えない（解析時の入力ごと保持）"),
    (1, "表の一括貼り付け・Delete クリア・まとめて Undo"),
    (1, ".json のドラッグ＆ドロップ、コマンドラインからの起動"),
    (1, "群杭沈下の入力を独立ウィンドウ (F8) に分離"),
]
col2 = [
    (0, "計算書 (docx)"),
    (1, "検定結果を Word の表に。主要な表に表番号と表題"),
    (1, "杭のモデル図を杭ごとに出力。印刷して読める文字サイズに作り直し"),
    (1, "「計算条件・仮定」章を新設"),
    (1, "　単位系・符号規約・選んだ材料オプションを明記"),
    (1, "解析後に条件を変えて出力した場合は、注意書きを表示"),
    (0, "速さ"),
    (1, "水平解析 76% 短縮（2,448 秒 → 588 秒）。荷重ケースは並列実行"),
    (1, "計算書の出力 5 倍（13 秒 → 2.6 秒）、Undo 8〜9 倍"),
]
bullets(s, Inches(0.72), Inches(1.55), Inches(5.5), col1, size=14.5, gap=14)
bullets(s, Inches(6.85), Inches(1.55), Inches(5.8), col2, size=14.5, gap=14)
picture(s, "report_page.png", Inches(6.9), Inches(5.3), Inches(1.55))
caption(s, Inches(8.75), Inches(6.02), Inches(3.6),
        "計算書のページ（図と表を作り直しました）", align=PP_ALIGN.LEFT)

# ── 6. 確からしさ ────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
header(s, 6, "確からしさをどう担保しているか",
       "結果が変わる修正も含むため、検証の仕組みを同時に増やしました。")

items = [
    (0, "文献との照合"),
    (1, "文献に数値として書かれた 45 項目を計算し、45 項目すべてが許容内"),
    (1, "出典と値を 1 か所に持ち、テストが照合したうえで検証ウィンドウと README の\n"
        "　  表を自動生成。表とテストが食い違わないようにしています"),
    (0, "自動テスト"),
    (1, "テスト 295 件 → 約 1,500 件（約 5 倍）。GitHub Actions で毎回自動実行"),
    (1, "「ビルドは通るが実行時に静かに壊れる」種類の不具合を、画面定義・コマンド・\n"
        "　  ヘルプ・用語・メッセージまで機械的に検査"),
    (1, "同じ入力なら毎回まったく同じ結果になることを検証"),
    (0, "不具合の修正（本資料では詳細を割愛）"),
    (1, "結果が変わる修正を含め継続的に対応。全件を CHANGELOG.md に、\n"
        "　  利用者向けの概要をヘルプの更新履歴に記録"),
]
bullets(s, Inches(0.72), Inches(1.52), Inches(11.9), items, size=14)

rect(s, Inches(0.72), Inches(6.1), SW - Inches(1.44), Inches(0.8), LIGHT)
t = tb(s, Inches(1.05), Inches(6.2), SW - Inches(2.1), Inches(0.6), anchor=MSO_ANCHOR.MIDDLE)
para(t, "この先 ─ 計算書の目次と総括表、解析の進行表示、製品名まわりの整理を予定",
     size=14.5, bold=True, color=NAVY, first=True)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print("saved:", OUT)
print(os.path.getsize(OUT), "bytes,", len(prs.slides._sldIdLst), "slides")
